from pathlib import Path
import csv
import re
import zipfile
import warnings
from collections import Counter, defaultdict

import pandas as pd
from openpyxl import load_workbook


# =============================================================================
# 07c - AUDIT PVSYST EXPORT METADATA / PROVENANCE
# =============================================================================
#
# PURPOSE
# -------
# Compare metadata and workbook structure across all six annual PVsyst exports:
#
#   A/B/C x Perez-Driesse / Engerer2
#
# with special attention to:
#
#   PVsyst_3_E2_h.xlsx   (System A / Engerer2)
#
# because scripts 07a/07b found a strong +1 h content offset in GTI,
# temperature and power only for this one export.
#
# This script is DIAGNOSTIC ONLY.
# It does not modify any source file.
#
# WHAT IT CHECKS
# --------------
# 1. Workbook/sheet structure
# 2. First 30 rows of every sheet
# 3. Keyword-bearing metadata lines such as:
#       meteo, weather, climate, project, variant, site, time,
#       hour, UTC, timezone, year, latitude, longitude, altitude
# 4. Workbook core properties
# 5. Defined names
# 6. Any filename/path-like strings visible in workbook XML, including
#       .MET, .CSV, .TXT, .SIT, .PRJ, .VCI, .PAN, .OND
# 7. Metadata lines unique to System A / Engerer2
#
# OUTPUT FOLDER
# -------------
# 03a_analysis_output_Engerer2/audit_SystemA_GTI/
#
# Outputs:
#   15_PVsyst_workbook_structure.csv
#   16_PVsyst_header_metadata.csv
#   17_PVsyst_keyword_hits.csv
#   18_PVsyst_defined_names.csv
#   19_PVsyst_embedded_file_references.csv
#   20_A_E2_unique_metadata.csv
#   21_A_E2_metadata_comparison.csv
#   PVsyst_EXPORT_METADATA_AUDIT.xlsx
#
# =============================================================================


ROOT = Path(__file__).resolve().parents[1]

SRC = ROOT / "01_source_inputs" / "annual" / "PVsyst"

OUT = (
    ROOT
    / "03a_analysis_output_Engerer2"
    / "audit_SystemA_GTI"
)
OUT.mkdir(parents=True, exist_ok=True)

FILES = {
    ("A", "PD"): SRC / "PVsyst_3_PD_h.xlsx",
    ("A", "E2"): SRC / "PVsyst_3_E2_h.xlsx",
    ("B", "PD"): SRC / "PVsyst_8_PD_h.xlsx",
    ("B", "E2"): SRC / "PVsyst_8_E2_h.xlsx",
    ("C", "PD"): SRC / "PVsyst_9_PD_h.xlsx",
    ("C", "E2"): SRC / "PVsyst_9_E2_h.xlsx",
}

TARGET = ("A", "E2")

HEADER_ROWS_TO_SCAN = 30
MAX_COLUMNS_TO_SCAN = 40

KEYWORDS = [
    "meteo",
    "meteorological",
    "weather",
    "climate",
    "project",
    "variant",
    "site",
    "location",
    "time",
    "hour",
    "timezone",
    "time zone",
    "utc",
    "gmt",
    "year",
    "date",
    "latitude",
    "longitude",
    "altitude",
    "azimuth",
    "tilt",
    "orientation",
    "simulation",
    "file",
    "source",
    "database",
    "station",
]

# File/path-like tokens that might identify the meteo input or PVsyst project.
FILE_TOKEN_RE = re.compile(
    r"""(?ix)
    (?:
        [A-Z]:\\[^\s<>"']+ |
        /[^\s<>"']+ |
        [^\s<>"']+
    )
    \.
    (?:met|csv|txt|sit|prj|vci|pan|ond|xlsx|xls|epw|tm2|tm3)
    """
)


# =============================================================================
# Helpers
# =============================================================================

def require_inputs():
    missing = [
        p for p in FILES.values()
        if not p.exists()
    ]

    if missing:
        raise FileNotFoundError(
            "Missing required PVsyst export(s):\n"
            + "\n".join(
                f"  - {p.relative_to(ROOT)}"
                for p in missing
            )
        )


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = re.sub(r"\s+", " ", text).strip()

    return text


def normalized_for_compare(value):
    """
    Normalize metadata text for cross-file comparison while preserving the
    original text in output tables.
    """
    text = normalize_text(value).lower()

    # Remove obvious changing whitespace only; do NOT strip numbers or paths.
    return text


def workbook_properties(wb):
    p = wb.properties

    fields = {
        "title": p.title,
        "subject": p.subject,
        "creator": p.creator,
        "keywords": p.keywords,
        "description": p.description,
        "lastModifiedBy": p.lastModifiedBy,
        "created": p.created,
        "modified": p.modified,
        "category": p.category,
        "contentStatus": p.contentStatus,
        "identifier": p.identifier,
        "language": p.language,
        "version": p.version,
    }

    return fields


def read_header_metadata(system, method, path):
    """
    Return:
      structure rows
      header cell rows
      keyword-hit rows
      defined-name rows
      workbook property rows
    """
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=True,
    )

    structure_rows = []
    header_rows = []
    keyword_rows = []
    defined_rows = []
    property_rows = []

    # Core workbook properties.
    for key, value in workbook_properties(wb).items():
        if value is None:
            continue

        property_rows.append({
            "System": system,
            "Method": method,
            "file": path.name,
            "property": key,
            "value": normalize_text(value),
        })

    # Workbook defined names.
    try:
        for name, defined_name in wb.defined_names.items():
            defined_rows.append({
                "System": system,
                "Method": method,
                "file": path.name,
                "defined_name": name,
                "value": normalize_text(
                    getattr(
                        defined_name,
                        "attr_text",
                        defined_name,
                    )
                ),
            })
    except Exception as exc:
        defined_rows.append({
            "System": system,
            "Method": method,
            "file": path.name,
            "defined_name": "<read error>",
            "value": str(exc),
        })

    for sheet_index, ws in enumerate(wb.worksheets):
        structure_rows.append({
            "System": system,
            "Method": method,
            "file": path.name,
            "sheet_index": sheet_index,
            "sheet_name": ws.title,
            "sheet_state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
        })

        max_row = min(
            HEADER_ROWS_TO_SCAN,
            ws.max_row or HEADER_ROWS_TO_SCAN,
        )
        max_col = min(
            MAX_COLUMNS_TO_SCAN,
            ws.max_column or MAX_COLUMNS_TO_SCAN,
        )

        for r in range(1, max_row + 1):
            row_values = []

            for c in range(1, max_col + 1):
                value = ws.cell(
                    row=r,
                    column=c,
                ).value

                text = normalize_text(value)

                if not text:
                    continue

                row_values.append(
                    f"{ws.cell(row=r, column=c).coordinate}={text}"
                )

                header_rows.append({
                    "System": system,
                    "Method": method,
                    "file": path.name,
                    "sheet_name": ws.title,
                    "row": r,
                    "column": c,
                    "cell": ws.cell(
                        row=r,
                        column=c,
                    ).coordinate,
                    "value": text,
                    "normalized_value": normalized_for_compare(text),
                })

                lower = text.lower()

                hits = [
                    kw for kw in KEYWORDS
                    if kw in lower
                ]

                if hits:
                    keyword_rows.append({
                        "System": system,
                        "Method": method,
                        "file": path.name,
                        "sheet_name": ws.title,
                        "row": r,
                        "cell": ws.cell(
                            row=r,
                            column=c,
                        ).coordinate,
                        "value": text,
                        "keywords": " | ".join(hits),
                    })

            # Also create one combined row-level record because metadata often
            # spans several cells on the same row.
            if row_values:
                combined = " || ".join(row_values)
                lower = combined.lower()

                hits = [
                    kw for kw in KEYWORDS
                    if kw in lower
                ]

                if hits:
                    keyword_rows.append({
                        "System": system,
                        "Method": method,
                        "file": path.name,
                        "sheet_name": ws.title,
                        "row": r,
                        "cell": "<combined row>",
                        "value": combined,
                        "keywords": " | ".join(hits),
                    })

    wb.close()

    return (
        structure_rows,
        header_rows,
        keyword_rows,
        defined_rows,
        property_rows,
    )


def embedded_file_references(system, method, path):
    """
    Search all XML/text parts inside the .xlsx zip package for filename/path
    references. This can reveal source paths not visible in worksheet cells.
    """
    rows = []
    seen = set()

    with zipfile.ZipFile(path, "r") as zf:
        for member in zf.namelist():
            # Only search text-like package parts.
            if not member.lower().endswith(
                (
                    ".xml",
                    ".rels",
                    ".txt",
                    ".vml",
                )
            ):
                continue

            try:
                data = zf.read(member)
                text = data.decode(
                    "utf-8",
                    errors="ignore",
                )
            except Exception:
                continue

            for match in FILE_TOKEN_RE.finditer(text):
                token = match.group(0).strip(
                    " ,;()[]{}"
                )

                key = (
                    member,
                    token.lower(),
                )

                if key in seen:
                    continue

                seen.add(key)

                rows.append({
                    "System": system,
                    "Method": method,
                    "file": path.name,
                    "package_part": member,
                    "reference": token,
                })

    return rows


# =============================================================================
# Cross-file comparison
# =============================================================================

def build_target_unique_metadata(
    header_df,
    keyword_df,
    property_df,
    defined_df,
    refs_df,
):
    """
    Extract metadata text present in A/E2 but not in any of the other five
    exports.
    """
    records = []

    sources = [
        (
            "header",
            header_df,
            "value",
        ),
        (
            "keyword",
            keyword_df,
            "value",
        ),
        (
            "property",
            property_df,
            "value",
        ),
        (
            "defined_name",
            defined_df,
            "value",
        ),
        (
            "embedded_reference",
            refs_df,
            "reference",
        ),
    ]

    for source_name, df, value_col in sources:
        if df.empty:
            continue

        d = df.copy()

        d["norm"] = d[value_col].map(
            normalized_for_compare
        )

        target = d[
            (d["System"] == TARGET[0])
            & (d["Method"] == TARGET[1])
        ].copy()

        others = d[
            ~(
                (d["System"] == TARGET[0])
                & (d["Method"] == TARGET[1])
            )
        ].copy()

        other_values = set(
            others["norm"]
            .dropna()
            .tolist()
        )

        target = target[
            ~target["norm"].isin(
                other_values
            )
        ]

        for _, r in target.iterrows():
            records.append({
                "source_type": source_name,
                "System": r["System"],
                "Method": r["Method"],
                "file": r["file"],
                "location": (
                    f"{r.get('sheet_name', '')}"
                    f":{r.get('cell', '')}"
                    if source_name in ("header", "keyword")
                    else r.get(
                        "package_part",
                        r.get(
                            "property",
                            r.get(
                                "defined_name",
                                "",
                            ),
                        ),
                    )
                ),
                "value": r[value_col],
            })

    return pd.DataFrame(records)


def build_metadata_comparison(keyword_df):
    """
    Put keyword-bearing lines side by side by normalized keyword concepts.
    This is intentionally simple and human-readable.
    """
    rows = []

    for kw in KEYWORDS:
        for (system, method), group in keyword_df.groupby(
            ["System", "Method"]
        ):
            hits = group[
                group["value"]
                .str.lower()
                .str.contains(
                    re.escape(kw),
                    regex=True,
                    na=False,
                )
            ]

            if hits.empty:
                continue

            combined = " || ".join(
                hits["value"]
                .drop_duplicates()
                .tolist()
            )

            rows.append({
                "keyword": kw,
                "System": system,
                "Method": method,
                "file": FILES[
                    (system, method)
                ].name,
                "matching_metadata": combined,
            })

    return pd.DataFrame(rows)


# =============================================================================
# Console summary
# =============================================================================

def print_console_summary(
    structure_df,
    keyword_df,
    refs_df,
    unique_df,
):
    print("\nWORKBOOK / SHEET STRUCTURE")
    print("-" * 80)

    print(
        structure_df[
            [
                "System",
                "Method",
                "file",
                "sheet_index",
                "sheet_name",
                "sheet_state",
                "max_row",
                "max_column",
            ]
        ].to_string(
            index=False
        )
    )

    print("\nKEY METADATA HITS FOR SYSTEM A")
    print("-" * 80)

    a_hits = keyword_df[
        keyword_df["System"] == "A"
    ].copy()

    show_cols = [
        "System",
        "Method",
        "file",
        "sheet_name",
        "row",
        "cell",
        "value",
        "keywords",
    ]

    if len(a_hits):
        print(
            a_hits[
                show_cols
            ].drop_duplicates().to_string(
                index=False,
                max_colwidth=100,
            )
        )
    else:
        print(
            "No keyword-bearing metadata found in the scanned worksheet header rows."
        )

    print("\nEMBEDDED FILE / PATH REFERENCES FOR SYSTEM A")
    print("-" * 80)

    a_refs = refs_df[
        refs_df["System"] == "A"
    ].copy()

    if len(a_refs):
        print(
            a_refs[
                [
                    "System",
                    "Method",
                    "file",
                    "package_part",
                    "reference",
                ]
            ].drop_duplicates().to_string(
                index=False,
                max_colwidth=100,
            )
        )
    else:
        print(
            "No recognizable file/path references found in workbook XML."
        )

    print("\nMETADATA UNIQUE TO A / ENGERER2")
    print("-" * 80)

    if len(unique_df):
        print(
            unique_df.to_string(
                index=False,
                max_colwidth=120,
            )
        )
    else:
        print(
            "No metadata text unique to A/E2 was found in the inspected workbook content."
        )

    print("\nINTERPRETATION")
    print("-" * 80)

    # Heuristic: prioritize unique references / time / meteo / project lines.
    if len(unique_df):
        suspicious = unique_df[
            unique_df["value"]
            .str.lower()
            .str.contains(
                r"meteo|weather|time|utc|gmt|project|variant|site|\.met|\.csv|\.txt",
                regex=True,
                na=False,
            )
        ]

        if len(suspicious):
            print(
                "A/E2 contains metadata/provenance entries not shared by the other exports "
                "that relate to meteo, time, project/site, or referenced files."
            )
            print(
                "Inspect 20_A_E2_unique_metadata.csv first. Any differing meteo/project "
                "reference could independently explain the one-hour content anomaly."
            )
        else:
            print(
                "A/E2 has some unique metadata, but nothing obviously identifies a "
                "different time convention or meteo source."
            )
            print(
                "If the source project/meteo file name is not present in the export, "
                "the next check must be made inside the A/E2 PVsyst project itself."
            )
    else:
        print(
            "The Excel exports do not expose an obvious A/E2-specific metadata difference."
        )
        print(
            "At that point, inspect the original A/E2 PVsyst project: imported meteo file, "
            "time convention, simulation variant, and export settings."
        )


# =============================================================================
# Main
# =============================================================================

def main():
    require_inputs()

    print(
        "\nPVSYST EXPORT METADATA / PROVENANCE AUDIT"
    )
    print("=" * 80)
    print(
        "Diagnostic only — no source or canonical files will be modified."
    )

    structure_rows = []
    header_rows = []
    keyword_rows = []
    defined_rows = []
    property_rows = []
    ref_rows = []

    for (system, method), path in FILES.items():
        (
            srows,
            hrows,
            krows,
            drows,
            prows,
        ) = read_header_metadata(
            system,
            method,
            path,
        )

        structure_rows.extend(
            srows
        )
        header_rows.extend(
            hrows
        )
        keyword_rows.extend(
            krows
        )
        defined_rows.extend(
            drows
        )
        property_rows.extend(
            prows
        )

        ref_rows.extend(
            embedded_file_references(
                system,
                method,
                path,
            )
        )

    structure_df = pd.DataFrame(
        structure_rows
    )
    header_df = pd.DataFrame(
        header_rows
    )
    keyword_df = pd.DataFrame(
        keyword_rows
    )
    defined_df = pd.DataFrame(
        defined_rows
    )
    property_df = pd.DataFrame(
        property_rows
    )
    refs_df = pd.DataFrame(
        ref_rows
    )

    # Ensure expected columns exist even if no rows were found.
    if refs_df.empty:
        refs_df = pd.DataFrame(
            columns=[
                "System",
                "Method",
                "file",
                "package_part",
                "reference",
            ]
        )

    if defined_df.empty:
        defined_df = pd.DataFrame(
            columns=[
                "System",
                "Method",
                "file",
                "defined_name",
                "value",
            ]
        )

    if property_df.empty:
        property_df = pd.DataFrame(
            columns=[
                "System",
                "Method",
                "file",
                "property",
                "value",
            ]
        )

    unique_df = build_target_unique_metadata(
        header_df,
        keyword_df,
        property_df,
        defined_df,
        refs_df,
    )

    comparison_df = build_metadata_comparison(
        keyword_df
    )

    outputs = {
        "15_PVsyst_workbook_structure.csv": structure_df,
        "16_PVsyst_header_metadata.csv": header_df,
        "17_PVsyst_keyword_hits.csv": keyword_df,
        "18_PVsyst_defined_names.csv": defined_df,
        "18b_PVsyst_workbook_properties.csv": property_df,
        "19_PVsyst_embedded_file_references.csv": refs_df,
        "20_A_E2_unique_metadata.csv": unique_df,
        "21_A_E2_metadata_comparison.csv": comparison_df,
    }

    print("\nOUTPUTS")
    print("-" * 80)

    for filename, df in outputs.items():
        path = OUT / filename

        df.to_csv(
            path,
            index=False,
        )

        print(
            f"Wrote {path.relative_to(ROOT)}"
        )

    # Optional Excel bundle.
    try:
        xlsx_path = (
            OUT
            / "PVsyst_EXPORT_METADATA_AUDIT.xlsx"
        )

        with pd.ExcelWriter(
            xlsx_path,
            engine="xlsxwriter",
        ) as writer:
            sheets = {
                "Workbook_structure": structure_df,
                "Header_metadata": header_df,
                "Keyword_hits": keyword_df,
                "Defined_names": defined_df,
                "Workbook_properties": property_df,
                "Embedded_refs": refs_df,
                "A_E2_unique": unique_df,
                "Keyword_comparison": comparison_df,
            }

            for name, df in sheets.items():
                df.to_excel(
                    writer,
                    sheet_name=name[:31],
                    index=False,
                )

                ws = writer.sheets[
                    name[:31]
                ]
                ws.freeze_panes(
                    1,
                    0,
                )

                for j, col in enumerate(
                    df.columns
                ):
                    values = [
                        len(str(v))
                        for v in df[col].head(200)
                        if pd.notna(v)
                    ]

                    width = max(
                        [len(str(col))]
                        + values
                    )

                    ws.set_column(
                        j,
                        j,
                        min(
                            max(width + 2, 10),
                            45,
                        ),
                    )

        print(
            f"Wrote {xlsx_path.relative_to(ROOT)}"
        )

    except Exception as exc:
        warnings.warn(
            f"Optional Excel bundle not created: {exc}"
        )

    print_console_summary(
        structure_df,
        keyword_df,
        refs_df,
        unique_df,
    )

    print("\nDONE")
    print("=" * 80)
    print(
        "If the export contains no useful meteo/project provenance, the next "
        "independent check must be performed in the original System A Engerer2 "
        "PVsyst project or its imported meteo file."
    )


if __name__ == "__main__":
    main()
